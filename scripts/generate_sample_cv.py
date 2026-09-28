"""Generate the fictional sample CV used by the browser E2E tests.

The candidate ("Alex Example") and every detail in it are invented test data: the real master
CV never enters the repository. The layout mixes Word heading styles, bold "caps" headings,
bullets and date ranges so the parser is exercised the same way as by real documents.

Usage (from the repository root):
    uv run python scripts/generate_sample_cv.py [OUTPUT_PATH]
"""

from __future__ import annotations

import sys
from datetime import UTC, datetime
from pathlib import Path

import docx

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = REPO_ROOT / "playwright" / "fixtures" / "sample-cv.docx"

# (style, text): title, heading (Word "Heading 1"), caps (bold paragraph), bold, bullet, text
SAMPLE_CV: tuple[tuple[str, str], ...] = (
    ("title", "Alex Example"),
    ("text", "AI Engineer"),
    ("text", "alex.example@example.com | +33 6 12 34 56 78 | linkedin.com/in/alex-example"),
    ("heading", "Summary"),
    ("text", "AI engineer building LLM and RAG systems in production."),
    ("caps", "EXPERIENCE"),
    ("bold", "Senior AI Engineer — Acme Analytics | Paris, France"),
    ("text", "Jan 2022 – Present"),
    ("bullet", "Designed a RAG platform with LangGraph and a vector database serving 2,000 users."),
    ("bullet", "Deployed models on Kubernetes (k8s) with Docker and FastAPI."),
    ("bold", "Machine Learning Engineer — Beta Labs"),
    ("text", "Sep 2019 – Dec 2021"),
    ("bullet", "Built time-series forecasting models with Spark."),
    ("bullet", "Reduced inference latency by 35%."),
    ("heading", "Education"),
    ("bold", "MSc in Computer Science — Université de Tunis"),
    ("text", "2017 – 2019"),
    ("heading", "Projects"),
    ("bold", "Job Agent — personal project"),
    ("bullet", "Multi-agent system orchestrating LLMs with MCP."),
    ("heading", "Skills"),
    ("text", "Languages: Python, SQL"),
    ("text", "ML: PyTorch, Scikit-learn, LLMs, RAG"),
    ("text", "Cloud: Azure, GCP"),
    ("heading", "Certifications"),
    ("bullet", "Azure AI Engineer Associate (2023)"),
    ("heading", "Languages"),
    ("text", "Arabic (native), French (fluent), English (professional)"),
)


def build(lines: tuple[tuple[str, str], ...]) -> docx.document.Document:
    document = docx.Document()
    document.core_properties.author = "Job Agent test fixture"
    document.core_properties.created = datetime(2026, 1, 1, tzinfo=UTC)
    for style, text in lines:
        if style == "title":
            document.add_paragraph(text, style="Title")
        elif style == "heading":
            document.add_heading(text, level=1)
        elif style == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif style in ("bold", "caps"):
            document.add_paragraph().add_run(text).bold = True
        else:
            document.add_paragraph(text)
    return document


def main() -> None:
    output = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUTPUT
    output.parent.mkdir(parents=True, exist_ok=True)
    build(SAMPLE_CV).save(str(output))
    print(f"Sample CV written to {output}")


if __name__ == "__main__":
    main()
