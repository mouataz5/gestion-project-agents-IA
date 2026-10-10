"""Backend test fixtures: synthetic CV documents, candidate profile files, jobs and discovery.

All CV content is fictional ("Alex Example"); the real master CV never enters the test suite.
Documents are generated at test time (python-docx / reportlab), so no binary fixtures are
committed.
"""

from __future__ import annotations

import io
import json
import shutil
import uuid
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import REPO_ROOT, Settings
from app.core.storage import LocalStorageProvider
from app.crawlers.base import CompanyTarget, JobQuery
from app.crawlers.mock import MockAtsSource
from app.cv.files import CvFileKind
from app.cv.models import ParsedCV
from app.cv.parser import parse_cv
from app.db.session import create_db_engine, create_session_factory
from app.llm.base import LLMProvider
from app.main import create_app
from app.models import Job, RunTrigger, RunType
from app.services.analysis import AnalysisService
from app.services.candidate_profile import CandidateService
from app.services.companies import CompanyService
from app.services.discovery import DiscoveryService
from app.services.master_cv import MasterCvService
from app.services.runs import RunService
from app.services.tailoring import CvTailoringService

# (style, text) — styles: title, heading (Word heading style), caps (plain bold paragraph used
# as a heading), bold, bullet, text.
CvLines = Sequence[tuple[str, str]]

SAMPLE_CV_EN: CvLines = [
    ("title", "Alex Example"),
    ("text", "AI Engineer"),
    ("text", "alex.example@example.com | +33 6 12 34 56 78 | linkedin.com/in/alex-example"),
    ("heading", "Summary"),
    ("text", "AI engineer building LLM and RAG systems in production."),
    ("caps", "EXPERIENCE"),
    ("bold", "Senior AI Engineer — Acme Analytics | Paris, France"),
    ("text", "Jan 2022 – Present"),
    ("bullet", "Designed a RAG platform with LangGraph and a vector database serving 2,000 users."),
    ("bullet", "Deployed models on Kubernetes (k8s) with Docker and FastAPI."),
    ("bold", "Machine Learning Engineer — Beta Labs"),
    ("text", "Sep 2019 – Dec 2021"),
    ("bullet", "Built time-series forecasting models with Spark."),
    ("bullet", "Reduced inference latency by 35%."),
    ("heading", "Education"),
    ("bold", "MSc in Computer Science — Université de Tunis"),
    ("text", "2017 – 2019"),
    ("heading", "Projects"),
    ("bold", "Job Agent — personal project"),
    ("bullet", "Multi-agent system orchestrating LLMs with MCP."),
    ("heading", "Skills"),
    ("text", "Languages: Python, SQL"),
    ("text", "ML: PyTorch, Scikit-learn, LLMs, RAG"),
    ("text", "Cloud: Azure, GCP"),
    ("heading", "Certifications"),
    ("bullet", "Azure AI Engineer Associate (2023)"),
    ("heading", "Languages"),
    ("text", "Arabic (native), French (fluent), English (professional)"),
    ("heading", "Interests"),
    ("text", "Chess, running"),
]

SAMPLE_CV_FR: CvLines = [
    ("title", "Alex Exemple"),
    ("text", "Ingénieur IA"),
    ("caps", "PROFIL"),
    ("text", "Ingénieur IA spécialisé en IA générative."),
    ("caps", "EXPÉRIENCE PROFESSIONNELLE"),
    ("bold", "Ingénieur IA — Société Démo"),
    ("text", "janv. 2021 – aujourd'hui"),
    ("bullet", "Conception d'agents LLM avec LangGraph."),
    ("bold", "Data Scientist — Startup Exemple"),
    ("text", "mars 2019 – déc. 2020"),
    ("bullet", "Modèles de vision par ordinateur."),
    ("caps", "FORMATION"),
    ("bold", "Diplôme d'ingénieur en informatique — ENSI"),
    ("text", "2014 – 2017"),
    ("caps", "COMPÉTENCES"),
    ("text", "Python, PyTorch, Docker"),
    ("caps", "LANGUES"),
    ("text", "Arabe, Français, Anglais"),
]


def build_docx(lines: CvLines) -> bytes:
    document = docx.Document()
    for style, text in lines:
        if style == "title":
            document.add_paragraph(text, style="Title")
        elif style == "heading":
            document.add_heading(text, level=1)
        elif style == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif style in ("bold", "caps"):
            document.add_paragraph().add_run(text).bold = True
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(lines: CvLines, *, pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for page in range(pages):
        y = 800.0
        for style, text in lines if page == 0 else [("text", f"Page {page + 1}")]:
            if style == "title":
                pdf.setFont("Helvetica-Bold", 16)
            elif style in ("heading", "caps", "bold"):
                pdf.setFont("Helvetica-Bold", 11)
            else:
                pdf.setFont("Helvetica", 10)
            pdf.drawString(72, y, f"- {text}" if style == "bullet" else text)
            y -= 18
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


@pytest.fixture
def cv_docx() -> Callable[[CvLines], bytes]:
    return build_docx


@pytest.fixture
def cv_pdf() -> Callable[..., bytes]:
    return build_pdf


@pytest.fixture
def sample_cv_en() -> CvLines:
    return SAMPLE_CV_EN


@pytest.fixture
def sample_cv_fr() -> CvLines:
    return SAMPLE_CV_FR


@pytest.fixture
def sample_master_cv() -> ParsedCV:
    """``SAMPLE_CV_EN`` as the real parser reads it (the structure a confirmed master CV holds)."""
    cv, _text = parse_cv(build_docx(SAMPLE_CV_EN), CvFileKind.DOCX)
    return cv


MOCK_JOBS_DIR = REPO_ROOT / "crawler" / "fixtures" / "mock_jobs"
JOB_FIELDS = (
    "company",
    "title",
    "description",
    "location",
    "country",
    "country_code",
    "remote_status",
    "employment_type",
    "seniority",
    "application_url",
    "ats_type",
    "visa_information",
    "relocation_information",
    "required_skills",
    "preferred_skills",
    "languages",
    "education_requirements",
    "experience_requirements",
    "responsibilities",
    "content_hash",
)


@pytest.fixture
def mock_board_ids() -> list[tuple[str, str]]:
    """Every (board token, job id) of the mock ATS boards."""
    pairs: list[tuple[str, str]] = []
    for path in sorted((MOCK_JOBS_DIR / "boards").glob("*.json")):
        board = json.loads(path.read_text(encoding="utf-8"))
        pairs.extend((path.stem, str(job["id"])) for job in board.get("jobs", []))
    return pairs


@pytest.fixture
def board_job() -> Callable[[str, str], Job]:
    """A posting of a mock ATS board, normalized exactly as discovery stores it (not persisted)."""

    def _job(board: str, source_job_id: str) -> Job:
        source = MockAtsSource(MOCK_JOBS_DIR)
        company = CompanyTarget(
            id=uuid.uuid4(), name=board, ats_type="GREENHOUSE", board_token=board, career_url=None
        )
        now = datetime(2026, 10, 6, 8, 0, tzinfo=UTC)
        query = JobQuery(
            posted_after=now - timedelta(days=30),
            posted_before=now,
            now=now,
            companies=(company,),
        )
        raw = next(item for item in source.search(query) if item.source_job_id == source_job_id)
        posting = source.normalize(raw)
        return Job(
            source=posting.source,
            source_job_id=posting.source_job_id,
            **{name: getattr(posting, name) for name in JOB_FIELDS},
        )

    return _job


@pytest.fixture
def candidate_dir(tmp_path: Path) -> Path:
    """A candidate folder with a copy of the repository's profile.yaml."""
    directory = tmp_path / "candidate"
    directory.mkdir()
    shutil.copy(REPO_ROOT / "candidate" / "profile.yaml", directory / "profile.yaml")
    return directory


@pytest.fixture
def candidate_settings(
    make_settings: Callable[..., Settings],
    postgres_url: str,
    clean_db: None,
    candidate_dir: Path,
) -> Settings:
    return make_settings(database_url=postgres_url, candidate_dir=candidate_dir)


@pytest.fixture
def candidate_client(
    candidate_settings: Settings, recording_queue: object, api_token: str
) -> Iterator[TestClient]:
    app = create_app(candidate_settings, task_queue=recording_queue)  # type: ignore[arg-type]
    with TestClient(app, headers={"Authorization": f"Bearer {api_token}"}) as client:
        yield client


# ---------------------------------------------------------------------------------------------
# Jobs and discovery (Phase 3)
# ---------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class RunSnapshot:
    id: uuid.UUID
    status: str
    jobs_discovered: int
    jobs_processed: int
    error_count: int
    summary: dict[str, Any]
    events: list[tuple[str, str, str]]  # (stage, level, message)
    jobs_qualified: int = 0
    cv_generated: int = 0


def snapshot_run(session_factory: sessionmaker[Session], run_id: uuid.UUID) -> RunSnapshot:
    with session_factory() as session:
        run = RunService(session).get_run(run_id)
        return RunSnapshot(
            id=run.id,
            status=run.status.value,
            jobs_discovered=run.jobs_discovered,
            jobs_processed=run.jobs_processed,
            error_count=run.error_count,
            summary=dict(run.summary),
            events=[(e.stage, e.level.value, e.message) for e in run.events],
            jobs_qualified=run.jobs_qualified,
            cv_generated=run.cv_generated,
        )


@pytest.fixture
def jobs_settings(
    make_settings: Callable[..., Settings],
    postgres_url: str,
    clean_db: None,
    candidate_dir: Path,
) -> Settings:
    return make_settings(database_url=postgres_url, candidate_dir=candidate_dir)


@pytest.fixture
def jobs_db(jobs_settings: Settings) -> Iterator[sessionmaker[Session]]:
    engine = create_db_engine(jobs_settings)
    yield create_session_factory(engine)
    engine.dispose()


@pytest.fixture
def import_companies(jobs_db: sessionmaker[Session]) -> Callable[[], None]:
    def _import() -> None:
        with jobs_db() as session:
            CompanyService(session).import_from_yaml(REPO_ROOT / "crawler")
            session.commit()

    return _import


@pytest.fixture
def discover(jobs_settings: Settings, jobs_db: sessionmaker[Session]) -> Callable[..., RunSnapshot]:
    """Run a discovery synchronously (as the worker does) and return a snapshot of the run."""

    def _run(*, clock: Callable[[], datetime] | None = None, **overrides: Any) -> RunSnapshot:
        settings = jobs_settings.model_copy(update=overrides) if overrides else jobs_settings
        with jobs_db() as session:
            run = RunService(session).create_run(
                run_type=RunType.DISCOVERY, trigger=RunTrigger.MANUAL
            )
            session.commit()
            run_id = run.id
        DiscoveryService(settings=settings, session_factory=jobs_db, clock=clock).execute(run_id)
        return snapshot_run(jobs_db, run_id)

    return _run


@pytest.fixture
def confirm_master_cv(
    jobs_settings: Settings,
    jobs_db: sessionmaker[Session],
    cv_docx: Callable[[CvLines], bytes],
    sample_cv_en: CvLines,
    tmp_path: Path,
) -> Callable[[], uuid.UUID]:
    """Upload and confirm the fictional sample CV for the default candidate."""

    def _confirm() -> uuid.UUID:
        with jobs_db() as session:
            candidate, _ = CandidateService(
                session, candidate_dir=jobs_settings.candidate_dir
            ).get_or_import_default()
            cvs = MasterCvService(session, LocalStorageProvider(tmp_path / "cv-storage"))
            version = cvs.upload(candidate, filename="cv.docx", data=cv_docx(sample_cv_en))
            cvs.confirm(candidate, version.id)
            session.commit()
            return version.id

    return _confirm


@pytest.fixture
def analyse(jobs_settings: Settings, jobs_db: sessionmaker[Session]) -> Callable[..., RunSnapshot]:
    """Run a job analysis synchronously (as the worker does) and return a snapshot of the run."""

    def _run(
        *,
        provider: LLMProvider | None = None,
        job_ids: Sequence[uuid.UUID] | None = None,
        force: bool = False,
        **overrides: Any,
    ) -> RunSnapshot:
        settings = jobs_settings.model_copy(update=overrides) if overrides else jobs_settings
        parameters: dict[str, Any] = {"force": force}
        if job_ids is not None:
            parameters["job_ids"] = [str(job_id) for job_id in job_ids]
        with jobs_db() as session:
            run = RunService(session).create_run(
                run_type=RunType.ANALYSIS, trigger=RunTrigger.MANUAL, parameters=parameters
            )
            session.commit()
            run_id = run.id
        AnalysisService(settings=settings, session_factory=jobs_db, provider=provider).execute(
            run_id
        )
        return snapshot_run(jobs_db, run_id)

    return _run


# The tailoring clock: the sample CV's years of experience are counted on this day.
TAILORING_NOW = datetime(2026, 10, 6, 8, 0, tzinfo=UTC)


@pytest.fixture
def tailor(jobs_settings: Settings, jobs_db: sessionmaker[Session]) -> Callable[..., RunSnapshot]:
    """Run a CV generation synchronously (as the worker does) and return a snapshot of the run."""

    def _run(
        *,
        provider: LLMProvider | None = None,
        job_ids: Sequence[uuid.UUID] | None = None,
        force: bool = False,
        **overrides: Any,
    ) -> RunSnapshot:
        settings = jobs_settings.model_copy(update=overrides) if overrides else jobs_settings
        parameters: dict[str, Any] = {"force": force}
        if job_ids is not None:
            parameters["job_ids"] = [str(job_id) for job_id in job_ids]
        with jobs_db() as session:
            run = RunService(session).create_run(
                run_type=RunType.CV_GENERATION, trigger=RunTrigger.MANUAL, parameters=parameters
            )
            session.commit()
            run_id = run.id
        CvTailoringService(
            settings=settings,
            session_factory=jobs_db,
            provider=provider,
            clock=lambda: TAILORING_NOW,
        ).execute(run_id)
        return snapshot_run(jobs_db, run_id)

    return _run


@pytest.fixture
def jobs_client(
    jobs_settings: Settings, recording_queue: object, api_token: str
) -> Iterator[TestClient]:
    app = create_app(jobs_settings, task_queue=recording_queue)  # type: ignore[arg-type]
    with TestClient(app, headers={"Authorization": f"Bearer {api_token}"}) as client:
        yield client
